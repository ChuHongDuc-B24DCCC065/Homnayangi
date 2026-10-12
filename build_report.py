# -*- coding: utf-8 -*-
"""
Script to generate the comprehensive, professional software engineering project report
for the Android application "Hôm Nay Ăn Gì?" (com.example.homnayangi).
Output: BAO_CAO_DU_AN_HOMNAYANGI.docx
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BÁO CÁO DỰ ÁN PHẦN MỀM: HÔM NAY ĂN GÌ? (ANDROID NATIVE)")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Dự án mã nguồn mở Android - Báo cáo kỹ thuật toàn diện")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

def format_paragraph(p, space_before=3, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_p(doc, text="", bold=False, italic=False, font_size=11, color=None, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    format_paragraph(p, space_before=2, space_after=space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = color
    return p

def add_callout(doc, title, text, bg_hex="FFF5F0", border_color_hex="FF6B35"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(190, 60, 20)
    
    run_c = p.add_run(text)
    run_c.font.name = "Arial"
    run_c.font.size = Pt(10)
    run_c.font.color.rgb = RGBColor(50, 50, 50)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_custom_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Arial"
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(230, 74, 0) # Deep Orange
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = RGBColor(33, 33, 33)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11.5)
        run.bold = True
        run.font.color.rgb = RGBColor(70, 70, 70)
    return h

def add_code_block(doc, code_str, title=""):
    if title:
        p_t = doc.add_paragraph()
        format_paragraph(p_t, space_before=6, space_after=2)
        r_t = p_t.add_run(f"Mã nguồn tham chiếu: {title}")
        r_t.italic = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGBColor(100, 100, 100)
        
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="CCCCCC"/>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="E0E0E0"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="E0E0E0"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="E0E0E0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=2, space_after=2, line_spacing=1.0)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def create_table(doc, headers, data, col_widths=None, alignment=WD_TABLE_ALIGNMENT.CENTER):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = alignment
    table.autofit = False
    set_table_borders(table)
    
    # Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, header_text in enumerate(headers):
        cell = hdr_row.cells[i]
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        set_cell_background(cell, "FF6B35") # Primary brand orange
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, space_before=2, space_after=2)
        run = p.add_run(header_text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row = table.rows[row_idx + 1]
        bg = "FFFFFF" if row_idx % 2 == 0 else "FFF8F5"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=2, space_after=2)
            
            # Text alignment logic
            if col_idx == 0 and len(cell_value) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            run = p.add_run(str(cell_value))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(40, 40, 40)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

print("Helper functions defined successfully.")
